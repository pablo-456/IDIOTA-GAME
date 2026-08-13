#!/usr/bin/env python3
"""
Generador de Documentación Técnica IEEE (SDD) — IDIOTA-GAME
Produce: docs/IDIOTA-Documentacion-Tecnica-IEEE.docx
"""

from __future__ import annotations

import os
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUTPUT = Path(__file__).resolve().parent / "IDIOTA-Documentacion-Tecnica-IEEE.docx"
DOC_VERSION = "1.0"
DOC_DATE = date.today().strftime("%d de %B de %Y").replace(
    "January", "enero"
).replace("February", "febrero").replace("March", "marzo").replace(
    "April", "abril"
).replace("May", "mayo").replace("June", "junio").replace(
    "July", "julio"
).replace("August", "agosto").replace("September", "septiembre").replace(
    "October", "octubre"
).replace("November", "noviembre").replace("December", "diciembre")

FONT_NAME = "Times New Roman"
BODY_SIZE = Pt(10)
HEADING1_SIZE = Pt(12)
TITLE_SIZE = Pt(24)
SUBTITLE_SIZE = Pt(14)

ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X",
         "XI", "XII", "XIII", "XIV", "XV", "XVI", "XVII", "XVIII", "XIX", "XX"]

_table_counter = 0
_figure_counter = 0


def _set_run_font(run, size=BODY_SIZE, bold=False, italic=False, color=None):
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = size
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def _style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, line_spacing=1.15):
    paragraph.alignment = align
    pf = paragraph.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line_spacing


def add_paragraph(doc, text, *, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                  size=BODY_SIZE, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, size=size, bold=bold, italic=italic)
    _style_paragraph(p, align=align, space_after=space_after)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    size = HEADING1_SIZE if level == 1 else Pt(11)
    run = p.add_run(text)
    _set_run_font(run, size=size, bold=True)
    pf = p.paragraph_format
    pf.space_before = Pt(12 if level == 1 else 8)
    pf.space_after = Pt(6)
    pf.keep_with_next = True
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.25 * level)
    run = p.add_run(text)
    _set_run_font(run)
    _style_paragraph(p, space_after=3)
    return p


def add_code_block(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Courier New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    run.font.size = Pt(8)
    pf = p.paragraph_format
    pf.left_indent = Inches(0.3)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.0
    return p


def add_table(doc, headers, rows, caption=None):
    global _table_counter
    _table_counter += 1
    cap = caption or f"Tabla {ROMAN[_table_counter - 1] if _table_counter <= len(ROMAN) else _table_counter}"

    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        run = hdr_cells[i].paragraphs[0].add_run(h)
        _set_run_font(run, bold=True)

    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = ""
            run = cells[c_idx].paragraphs[0].add_run(str(val))
            _set_run_font(run)

    cap_p = doc.add_paragraph()
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap_p.add_run(cap)
    _set_run_font(run, italic=True, size=Pt(9))
    cap_p.paragraph_format.space_after = Pt(10)
    return table


def add_figure(doc, diagram_text, caption):
    global _figure_counter
    _figure_counter += 1
    add_code_block(doc, diagram_text)
    cap_p = doc.add_paragraph()
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap_p.add_run(f"Fig. {_figure_counter}. {caption}")
    _set_run_font(run, italic=True, size=Pt(9))
    cap_p.paragraph_format.space_after = Pt(10)


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_char_begin)

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    run._r.append(instr)

    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    run._r.append(fld_char_separate)

    run2 = p.add_run("Actualice la tabla de contenidos: clic derecho → Actualizar campo.")
    _set_run_font(run2, italic=True, size=Pt(9))

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_end)
    p.paragraph_format.space_after = Pt(12)


def collect_file_tree(root: Path, prefix="") -> list[str]:
    lines = []
    entries = sorted(root.iterdir(), key=lambda p: (not p.is_dir(), p.name.lower()))
    for i, entry in enumerate(entries):
        if entry.name in {".git", "node_modules", "__pycache__", ".cursor"}:
            continue
        connector = "└── " if i == len(entries) - 1 else "├── "
        lines.append(f"{prefix}{connector}{entry.name}{'/' if entry.is_dir() else ''}")
        if entry.is_dir() and entry.name not in {"node_modules", ".git", "__pycache__"}:
            extension = "    " if i == len(entries) - 1 else "│   "
            lines.extend(collect_file_tree(entry, prefix + extension))
    return lines


def build_cover(doc):
    for _ in range(4):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Documentación Técnica del Sistema IDIOTA")
    _set_run_font(run, size=TITLE_SIZE, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Juego de Cartas Multijugador en Tiempo Real")
    _set_run_font(run, size=SUBTITLE_SIZE, italic=True)

    doc.add_paragraph()
    meta_lines = [
        "Software Design Document (SDD)",
        f"Versión del documento: {DOC_VERSION}",
        f"Fecha: {DOC_DATE}",
        "Estado del sistema: Beta funcional",
        "URL de producción: https://idiota-game.vercel.app",
    ]
    for line in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        _set_run_font(run, size=Pt(11))

    doc.add_page_break()


def build_revision_table(doc):
    add_heading(doc, "Control de revisiones", level=1)
    add_table(
        doc,
        ["Versión", "Fecha", "Descripción"],
        [
            ["1.0", DOC_DATE, "Documento inicial: arquitectura, diseño backend/frontend, contratos Socket.io y reglas de negocio."],
        ],
        caption="Tabla I. Control de revisiones del documento",
    )
    doc.add_page_break()


def build_document(doc):
    # ── 1. Introducción ──
    add_heading(doc, "1. Introducción")
    add_heading(doc, "1.1 Propósito", level=2)
    add_paragraph(
        doc,
        "El presente Software Design Document (SDD) describe la arquitectura, el diseño "
        "y la implementación del sistema IDIOTA, un juego de cartas multijugador online "
        "en tiempo real. Su objetivo es servir como referencia técnica para desarrolladores, "
        "mantenedores y evaluadores académicos que necesiten comprender, extender o auditar el proyecto."
    )

    add_heading(doc, "1.2 Alcance", level=2)
    add_paragraph(
        doc,
        "Este documento cubre el mono-repositorio IDIOTA-GAME, compuesto por un backend "
        "Node.js con Socket.io y un frontend React empaquetado con Vite. Se documentan las "
        "reglas de negocio del juego, la máquina de estados, los contratos de comunicación "
        "WebSocket, la estructura de componentes del cliente y el despliegue en producción. "
        "Quedan fuera del alcance: persistencia en base de datos, autenticación de usuarios "
        "y sistemas de pago o monetización."
    )

    add_heading(doc, "1.3 Audiencia", level=2)
    add_bullet(doc, "Desarrolladores que mantengan o extiendan el código fuente.")
    add_bullet(doc, "Revisores técnicos y evaluadores académicos.")
    add_bullet(doc, "Operadores que desplieguen el sistema en entornos de prueba o producción.")

    add_heading(doc, "1.4 Resumen ejecutivo", level=2)
    add_paragraph(
        doc,
        "IDIOTA es un juego de cartas para 3 a 7 jugadores conectados simultáneamente "
        "mediante WebSockets. El servidor mantiene todo el estado de las partidas en memoria "
        "RAM, valida las jugadas de forma autoritativa y sincroniza el estado con los clientes "
        "web. El último jugador que conserva cartas al final de la partida es declarado "
        "\"idiota\" (perdedor); los demás jugadores que se quedan sin cartas son \"salvados\"."
    )

    # ── 2. Referencias ──
    add_heading(doc, "2. Referencias")
    refs = [
        "[1] README del proyecto — IDIOTA-GAME/README.md",
        "[2] Reglas internas del repositorio — .cursor/rules/ (01–06)",
        "[3] Dependencias backend — backend/package.json",
        "[4] Dependencias frontend — frontend/package.json",
        "[5] Documentación oficial de Socket.io — https://socket.io/docs/v4/",
        "[6] Documentación oficial de React — https://react.dev/",
        "[7] Documentación oficial de Vite — https://vitejs.dev/",
        "[8] Estándar IEEE 1016 — Software Design Descriptions",
    ]
    for ref in refs:
        add_bullet(doc, ref)

    # ── 3. Definiciones ──
    add_heading(doc, "3. Definiciones, acrónimos y abreviaturas")
    add_table(
        doc,
        ["Término", "Definición"],
        [
            ["SDD", "Software Design Document — documento de diseño de software."],
            ["WebSocket", "Protocolo de comunicación full-duplex sobre TCP para tiempo real."],
            ["SRP", "Single Responsibility Principle — principio de responsabilidad única (SOLID)."],
            ["LOBBY", "Estado inicial de sala: esperando jugadores antes de comenzar."],
            ["SETUP", "Fase de configuración donde cada jugador elige cartas visibles y mano."],
            ["PLAYING", "Estado de partida activa con turnos alternados."],
            ["FINISHED", "Estado terminal: partida concluida con perdedor identificado."],
            ["Idiota", "Jugador perdedor: el último que conserva cartas al terminar."],
            ["Saved / Salvado", "Jugador que se queda sin cartas antes del final; pasa a modo espectador."],
            ["Pile / Pozo", "Pila central de cartas jugadas sobre la mesa."],
            ["Host", "Primer jugador de la sala; único autorizado a iniciar el setup."],
        ],
        caption="Tabla II. Glosario de términos y acrónimos",
    )

    # ── 4. Descripción general ──
    add_heading(doc, "4. Descripción general del sistema")

    add_heading(doc, "4.1 Perspectiva del producto", level=2)
    add_paragraph(
        doc,
        "IDIOTA opera como una aplicación web cliente-servidor. El cliente (React/Vite) "
        "se ejecuta en el navegador del jugador y se comunica exclusivamente con el servidor "
        "mediante Socket.io. No existe capa de persistencia: las salas y partidas viven "
        "únicamente en la memoria del proceso Node.js. Los jugadores se identifican por "
        "socket.id (asignado por Socket.io) y un nombre de usuario libre (username) sin "
        "autenticación ni registro."
    )

    add_heading(doc, "4.2 Funciones principales", level=2)
    functions = [
        "Creación de salas con código alfanumérico único de 6 caracteres.",
        "Unión a salas existentes mediante código compartido.",
        "Lobby interactivo con validación de mínimo 3 y máximo 7 jugadores.",
        "Fase de setup con elección estratégica de 4 cartas visibles y 4 en mano privada.",
        "Partida por turnos con validación de cartas, comodines y recogida de pozo.",
        "Efectos visuales (overlays) y sistema de audio en el cliente.",
        "Modo espectador para jugadores salvados.",
    ]
    for f in functions:
        add_bullet(doc, f)

    add_heading(doc, "4.3 Restricciones de diseño", level=2)
    add_bullet(doc, "Estado efímero en RAM: un reinicio del servidor elimina todas las salas activas.")
    add_bullet(doc, "Validación autoritativa en backend/src/models/Game.js; el cliente solo refleja el estado recibido.")
    add_bullet(doc, "CORS restringido a orígenes conocidos: Vercel, localhost y túneles devtunnels.ms.")
    add_bullet(doc, "Sin base de datos, caché externo ni cola de mensajes.")

    add_heading(doc, "4.4 Supuestos y dependencias", level=2)
    add_bullet(doc, "Node.js versión 18 o superior en el servidor.")
    add_bullet(doc, "Navegador moderno con soporte WebSocket en el cliente.")
    add_bullet(doc, "Variables de entorno: PORT, FRONTEND_URL (backend); VITE_SERVER_URL (frontend).")

    # ── 5. Arquitectura ──
    add_heading(doc, "5. Arquitectura del sistema")

    add_figure(
        doc,
        """┌─────────────────────────────────────────────────────────────┐
│                    FRONTEND (React + Vite)                     │
│  ┌──────────┐  ┌──────────────────────────────────────────┐ │
│  │  Pages   │  │  Home  →  Lobby  →  Game                 │ │
│  └────┬─────┘  └──────────────────────────────────────────┘ │
│       │                                                      │
│  ┌────▼──────────────────────────────────────────────────┐  │
│  │  Hooks: useSocket | useRoomSession | useGameActions   │  │
│  │         useGameEvents | useAudio                      │  │
│  └────┬──────────────────────────────────────────────────┘  │
│       │  Components: SetupPhase, PlayingPhase, Overlays    │
└───────┼──────────────────────────────────────────────────────┘
        │  Socket.io (WebSocket)
┌───────▼──────────────────────────────────────────────────────┐
│                    BACKEND (Node.js)                            │
│  ┌─────────────┐    ┌──────────────────┐    ┌─────────────┐  │
│  │  server.js  │───►│ roomController.js│───►│  Game.js    │  │
│  │  (Transporte)│    │  (Salas/RAM)     │    │  Deck.js    │  │
│  └─────────────┘    └──────────────────┘    │  (Dominio)  │  │
│                                              └─────────────┘  │
│  GET /health  —  Health check HTTP                             │
└────────────────────────────────────────────────────────────────┘""",
        "Arquitectura en capas del sistema IDIOTA",
    )

    add_heading(doc, "5.1 Estructura del mono-repositorio", level=2)
    add_code_block(
        doc,
        """IDIOTA-GAME/
├── backend/
│   ├── package.json
│   └── src/
│       ├── server.js              # Express + Socket.io
│       ├── controllers/
│       │   └── roomController.js  # Ciclo de vida de salas
│       └── models/
│           ├── Game.js            # Reglas y estado de partida
│           └── Deck.js            # Mazo de 108 cartas
├── frontend/
│   ├── package.json
│   └── src/
│       ├── App.jsx                # Enrutador por pantallas
│       ├── pages/                 # Home, Lobby, Game
│       ├── components/            # UI, overlays, cartas
│       ├── hooks/                 # Socket, sesión, acciones
│       ├── constants/             # Eventos y reglas de cartas
│       └── audio/                 # Motor Web Audio
└── docs/                          # Documentación generada""",
    )

    add_heading(doc, "5.2 Principios de diseño", level=2)
    add_paragraph(
        doc,
        "El proyecto aplica el principio de responsabilidad única (SRP): server.js gestiona "
        "únicamente transporte y emisión de eventos; Game.js concentra la lógica de negocio; "
        "roomController.js administra el ciclo de vida de salas. En el frontend, los hooks "
        "desacoplan la conexión de red de los componentes visuales, que permanecen "
        "presentacionales y reciben datos y callbacks como props."
    )

    # ── 6. Backend ──
    add_heading(doc, "6. Diseño detallado — Backend")

    add_heading(doc, "6.1 Máquina de estados", level=2)
    add_figure(
        doc,
        "LOBBY  ──start_setup──►  SETUP  ──allReady──►  PLAYING  ──gameOver──►  FINISHED",
        "Transiciones de estado de una partida IDIOTA",
    )
    add_table(
        doc,
        ["Estado", "Descripción", "Transición"],
        [
            ["LOBBY", "Sala abierta; jugadores se unen (3–7).", "Host ejecuta start_setup → SETUP"],
            ["SETUP", "Reparto y elección de cartas visibles/mano.", "Todos confirman → PLAYING"],
            ["PLAYING", "Turnos alternados; validación de jugadas.", "Último con cartas → FINISHED"],
            ["FINISHED", "Partida terminada; perdedor identificado.", "Sala se destruye al vaciarse"],
        ],
        caption="Tabla III. Estados de la partida",
    )

    add_heading(doc, "6.2 Modelo de datos del jugador", level=2)
    add_table(
        doc,
        ["Campo", "Tipo", "Descripción"],
        [
            ["id", "string", "socket.id del jugador (identificador de sesión)."],
            ["username", "string", "Nombre visible elegido al unirse."],
            ["manoPrivada", "Card[]", "Cartas en mano; ocultas para rivales."],
            ["cartasVisibles", "Card[]", "4 cartas boca arriba visibles para todos."],
            ["cartasOcultas", "Card[]", "4 cartas boca abajo; se usan al agotar mano y visibles."],
            ["isReady", "boolean", "Confirmó elección en fase SETUP."],
            ["saved", "boolean", "True si el jugador se quedó sin cartas (salvado)."],
        ],
        caption="Tabla IV. Estructura del objeto jugador",
    )

    add_heading(doc, "6.3 Mazo (Deck.js)", level=2)
    add_paragraph(
        doc,
        "El mazo consta de 108 cartas: dos barajas de póker completas (104 cartas) más "
        "4 comodines JOKER adicionales. La mezcla utiliza el algoritmo Fisher-Yates (Knuth Shuffle) "
        "con complejidad O(n), garantizando distribución uniforme."
    )
    add_table(
        doc,
        ["Tipo", "Valor", "Efecto"],
        [
            ["Ordinaria", "3–7, 9, 10, J, Q, K, A", "Jerarquía: 3(3) < … < A(14). Debe ser ≥ tope de pila."],
            ["Reset", "2", "Reinicia la pila; cualquier carta puede jugarse después."],
            ["Burn", "8", "Quema la pila completa; turno extra para quien la juega."],
            ["Burn", "JOKER", "Igual que el 8; 4 en total (2 por baraja)."],
        ],
        caption="Tabla V. Tipos de carta y efectos especiales",
    )

    add_heading(doc, "6.4 Lógica de turnos", level=2)
    add_bullet(doc, "El jugador activo juega una o más cartas del mismo valor cuyo power ≥ pileTopPower.")
    add_bullet(doc, "Prioridad de zonas: manoPrivada → cartasVisibles → cartasOcultas.")
    add_bullet(doc, "Regla Azar: carta oculta demasiado débil se revela; el jugador recoge el pozo tras 3 s.")
    add_bullet(doc, "Recogida voluntaria o forzosa: cartas del pozo pasan a manoPrivada; el turno NO avanza.")
    add_bullet(doc, "Cuatro cartas iguales consecutivas en la pila queman automáticamente la mesa.")
    add_bullet(doc, "Al final del turno, el jugador roba del mazo hasta tener 4 cartas en mano (si hay mazo).")

    add_heading(doc, "6.5 Gestión de salas (roomController.js)", level=2)
    add_paragraph(
        doc,
        "El controlador mantiene dos estructuras en memoria: games (roomId → instancia Game) "
        "y playerRoomMap (socketId → roomId) para búsqueda O(1) en desconexiones. "
        "Las salas se destruyen cuando quedan vacías o la partida termina sin jugadores suficientes."
    )

    add_heading(doc, "6.6 Endpoint HTTP", level=2)
    add_table(
        doc,
        ["Método", "Ruta", "Respuesta"],
        [
            ["GET", "/health", "{ status: 'ok', rooms: [...], uptime: number }"],
        ],
        caption="Tabla VI. Endpoints HTTP del backend",
    )

    # ── 7. Socket.io ──
    add_heading(doc, "7. Contrato de comunicación (Socket.io)")

    add_heading(doc, "7.1 Eventos cliente → servidor", level=2)
    add_table(
        doc,
        ["Evento", "Payload", "Validación / Efecto"],
        [
            ["create_room", "{ username }", "Crea sala; emite room_created y room_updated."],
            ["join_room", "{ roomId, username }", "Une a sala LOBBY; emite room_joined y room_updated."],
            ["start_setup", "{ roomId }", "Solo host; LOBBY → SETUP; emite setup_started por jugador."],
            ["confirm_setup", "{ roomId, idsVisibles[] }", "Exactamente 4 IDs visibles; si todos listos → PLAYING."],
            ["play_turn", "{ roomId, cardIds[] }", "Valida y ejecuta jugada; sincroniza con game_started."],
            ["pick_up_pile", "{ roomId, voluntary }", "Recoge pozo; sincroniza estado a todos."],
        ],
        caption="Tabla VII. Eventos Socket.io — Cliente → Servidor",
    )

    add_heading(doc, "7.2 Eventos servidor → cliente", level=2)
    add_table(
        doc,
        ["Evento", "Destinatario", "Payload principal"],
        [
            ["room_created", "Creador", "{ roomId, state } — estado privado del creador."],
            ["room_joined", "Nuevo jugador", "{ roomId, state } — estado privado del joiner."],
            ["room_updated", "Toda la sala", "Estado público (sin manos rivales)."],
            ["setup_started", "Cada jugador", "{ state } — cartas de elección privadas."],
            ["game_started", "Cada jugador", "{ state, firstPlayerId?, message? } — sync de juego."],
            ["card_revealed", "Toda la sala", "{ playerId, card, mustPickUp } — reveal dramático."],
            ["special_play", "Toda la sala", "{ playerId, card, burned } — animación 8/Joker."],
            ["player_saved", "Jugador salvado", "{ message, savedPlayers } — notificación personal."],
            ["player_disconnected", "Sala restante", "{ socketId, advanceTurn, state }."],
            ["game_over", "Toda la sala", "{ loserId, loserName, savedPlayers, reason }."],
            ["error", "Emisor", "{ message } — error de validación."],
        ],
        caption="Tabla VIII. Eventos Socket.io — Servidor → Cliente",
    )

    add_heading(doc, "7.3 Nota técnica de sincronización", level=2)
    add_paragraph(
        doc,
        "Durante la fase PLAYING, el servidor reutiliza el evento game_started como mecanismo "
        "de sincronización de estado completo (no existe un evento state_updated dedicado). "
        "Los clientes deben tratar game_started tanto como inicio de partida (cuando incluye "
        "firstPlayerId) como actualización intermedia de estado durante el juego."
    )

    add_heading(doc, "7.4 Diagrama de secuencia — Flujo feliz", level=2)
    add_figure(
        doc,
        """Jugador A          Servidor           Jugadores B..N
    │                    │                      │
    │── create_room ────►│                      │
    │◄── room_created ───│                      │
    │                    │                      │
    │                    │◄── join_room ────────│
    │                    │── room_joined ──────►│
    │◄── room_updated ───┼── room_updated ──────►│
    │                    │                      │
    │── start_setup ────►│                      │
    │◄── setup_started ──┼── setup_started ─────►│
    │                    │                      │
    │── confirm_setup ──►│◄── confirm_setup ────│
    │◄── game_started ───┼── game_started ─────►│
    │                    │                      │
    │── play_turn ──────►│                      │
    │◄── game_started ───┼── game_started ─────►│
    │       ...          │         ...          │
    │◄── game_over ──────┼── game_over ────────►│""",
        "Secuencia de eventos desde creación de sala hasta fin de partida",
    )

    # ── 8. Frontend ──
    add_heading(doc, "8. Diseño detallado — Frontend")

    add_heading(doc, "8.1 Navegación por pantallas", level=2)
    add_paragraph(
        doc,
        "El frontend no utiliza React Router. La navegación se controla mediante el estado "
        "screen del hook useRoomSession con tres valores: HOME, LOBBY y GAME. "
        "App.jsx renderiza condicionalmente Home, Lobby o Game según el valor actual."
    )
    add_figure(
        doc,
        "HOME ──create/join──► LOBBY ──setup_started──► GAME (SETUP)\n"
        "                                              │\n"
        "                                              ▼\n"
        "                                         GAME (PLAYING)\n"
        "                                              │\n"
        "                                              ▼\n"
        "                                         GAME (FINISHED)",
        "Flujo de navegación del cliente",
    )

    add_heading(doc, "8.2 Capas de hooks", level=2)
    add_table(
        doc,
        ["Hook", "Archivo", "Responsabilidad"],
        [
            ["useSocket", "hooks/useSocket.js", "Conexión singleton Socket.io; emit/on/off; detección URL prod/dev."],
            ["useRoomSession", "hooks/useRoomSession.js", "Estado de sesión, navegación HOME→LOBBY→GAME, gameState."],
            ["useGameActions", "hooks/useGameActions.js", "Emisión de eventos al servidor (createRoom, playTurn, etc.)."],
            ["useGameEvents", "hooks/useGameEvents.js", "Estado local de overlays: reveal, special, crown, saved."],
            ["useAudio", "hooks/useAudio.js", "Motor Web Audio, mute, SFX y música de fondo."],
        ],
        caption="Tabla IX. Hooks personalizados del frontend",
    )

    add_heading(doc, "8.3 Componentes principales", level=2)
    add_table(
        doc,
        ["Categoría", "Componente", "Función"],
        [
            ["Páginas", "Home", "Formulario username/código; crear o unirse a sala."],
            ["Páginas", "Lobby", "Lista de jugadores, código de sala, botón iniciar (host)."],
            ["Páginas", "Game", "Shell principal; compone fases y overlays."],
            ["Fases", "SetupPhase", "Animación de reparto y selección de 4 cartas visibles."],
            ["Fases", "PlayingPhase", "Mesa activa: mano, oponentes, pozo, acciones."],
            ["Fases", "SpectatorView", "Vista de solo lectura para jugadores salvados."],
            ["Cartas", "Card, PileDisplay", "Renderizado de carta individual y pila central."],
            ["Mesa", "BoardCenter, OpponentsRow, MyVisiblesPanel", "Layout de la mesa de juego."],
            ["Overlays", "DealOverlay, StartRouletteOverlay", "Animaciones de reparto y ruleta inicial."],
            ["Overlays", "CardRevealOverlay, SpecialPlayOverlay", "Efectos de reveal y cartas especiales."],
            ["Overlays", "CrownOverlay", "Revelación del perdedor al finalizar."],
            ["UI", "MuteButton", "Control flotante de silencio de audio."],
        ],
        caption="Tabla X. Componentes React del frontend",
    )

    add_heading(doc, "8.4 Identidad visual", level=2)
    add_paragraph(
        doc,
        "La interfaz sigue una estética de casino oscuro definida en frontend/src/styles/theme.css. "
        "Variables CSS principales: fondo verde fieltro (--felt), acentos dorados (--gold), "
        "texto crema (--cream). Tipografías: Playfair Display (títulos) y EB Garamond (cuerpo). "
        "Las cartas seleccionadas usan transform: translateY para efecto de elevación."
    )

    add_heading(doc, "8.5 Resolución de URL del servidor", level=2)
    add_paragraph(
        doc,
        "El hook useSocket detecta automáticamente la URL del backend en este orden: "
        "(1) variable VITE_SERVER_URL si está definida; (2) en desarrollo, localhost:3000 "
        "o túnel devtunnels con puerto 3000; (3) en producción, https://idiota-backend.onrender.com."
    )

    # ── 9. Reglas de negocio ──
    add_heading(doc, "9. Reglas de negocio del juego")

    add_heading(doc, "9.1 Reparto inicial (SETUP)", level=2)
    add_bullet(doc, "Cada jugador recibe 4 cartas ocultas (cartasOcultas) que no ve hasta usarlas.")
    add_bullet(doc, "Recibe 8 cartas adicionales para elegir: 4 quedan visibles, 4 en mano privada.")
    add_bullet(doc, "Todos deben confirmar antes de pasar a PLAYING.")
    add_bullet(doc, "Se elige al primer jugador al azar mediante ruleta visual en el cliente.")

    add_heading(doc, "9.2 Validación de jugadas", level=2)
    add_bullet(doc, "Las cartas jugadas deben ser del mismo valor.")
    add_bullet(doc, "El power de la carta más alta jugada debe ser ≥ pileTopPower (salvo pila vacía o tras reset).")
    add_bullet(doc, "El '2' (reset) puede jugarse sobre cualquier carta y deja la pila en power 0.")
    add_bullet(doc, "El '8' y JOKER queman la pila y otorgan turno extra.")

    add_heading(doc, "9.3 Reglas de recogida", level=2)
    add_bullet(doc, "Penalización: si no puede responder, recoge todo el pozo a su mano privada.")
    add_bullet(doc, "Recogida voluntaria: en su turno puede elegir recoger el pozo por estrategia.")
    add_bullet(doc, "En ambos casos el turno NO avanza; el mismo jugador abre la nueva ronda con pila vacía.")

    add_heading(doc, "9.4 Condición de victoria y derrota", level=2)
    add_paragraph(
        doc,
        "Un jugador se salva (saved) cuando se queda sin cartas en las tres zonas "
        "(mano, visibles y ocultas). La partida termina cuando queda un solo jugador "
        "con cartas: ese jugador es declarado idiota (perdedor). Mínimo 3 jugadores "
        "para iniciar; máximo 7 por sala."
    )

    # ── 10. Despliegue ──
    add_heading(doc, "10. Despliegue e infraestructura")
    add_table(
        doc,
        ["Componente", "Plataforma", "Configuración"],
        [
            ["Frontend", "Vercel", "Build: npm run build (Vite). Env: VITE_SERVER_URL."],
            ["Backend", "Render", "Start: node src/server.js. Env: PORT, FRONTEND_URL."],
            ["Desarrollo local", "localhost", "Backend :3000, Frontend :5173 (npm run dev)."],
        ],
        caption="Tabla XI. Entornos de despliegue",
    )

    add_heading(doc, "10.1 Instalación local", level=2)
    add_code_block(
        doc,
        """# Backend
cd backend
npm install
npm start          # Servidor en puerto 3000

# Frontend (terminal separada)
cd frontend
npm install
npm run dev        # Cliente en puerto 5173""",
    )

    add_heading(doc, "10.2 CORS y orígenes permitidos", level=2)
    add_bullet(doc, "process.env.FRONTEND_URL (configurable)")
    add_bullet(doc, "https://idiota-game.vercel.app")
    add_bullet(doc, "http://localhost:5173 y http://localhost:3000")
    add_bullet(doc, "Cualquier origen *.devtunnels.ms (desarrollo remoto)")

    # ── 11. Limitaciones ──
    add_heading(doc, "11. Limitaciones conocidas y trabajo futuro")
    add_table(
        doc,
        ["Limitación / Roadmap", "Estado"],
        [
            ["Sin persistencia de usuarios ni historial de partidas", "Actual"],
            ["Sin autenticación ni sistema anti-trampas avanzado", "Actual"],
            ["Salas públicas (matchmaking automático)", "Planificado — necesidades.txt"],
            ["Música de suspenso en fase final", "Planificado — necesidades.txt"],
        ],
        caption="Tabla XII. Limitaciones y hoja de ruta",
    )

    # ── 12. Apéndices ──
    add_heading(doc, "Apéndice A. Árbol de archivos del proyecto")
    tree_lines = collect_file_tree(ROOT)
    add_code_block(doc, "IDIOTA-GAME/\n" + "\n".join(tree_lines[:80]))
    if len(tree_lines) > 80:
        add_paragraph(doc, f"(… {len(tree_lines) - 80} entradas adicionales omitidas por brevedad)", italic=True)

    add_heading(doc, "Apéndice B. Dependencias del proyecto")
    add_paragraph(doc, "Backend (backend/package.json):", bold=True)
    add_table(
        doc,
        ["Paquete", "Versión", "Uso"],
        [
            ["express", "^4.22.2", "Servidor HTTP y endpoint /health."],
            ["socket.io", "^4.8.3", "Comunicación WebSocket bidireccional."],
            ["cors", "^2.8.6", "Declarado; CORS configurado en Socket.io."],
            ["nodemon", "^3.0.2", "Recarga en desarrollo (devDependency)."],
        ],
        caption="Tabla XIII. Dependencias del backend",
    )
    add_paragraph(doc, "Frontend (frontend/package.json):", bold=True)
    add_table(
        doc,
        ["Paquete", "Versión", "Uso"],
        [
            ["react", "^18.2.0", "Biblioteca de interfaz de usuario."],
            ["react-dom", "^18.2.0", "Renderizado DOM de React."],
            ["socket.io-client", "^4.7.2", "Cliente WebSocket."],
            ["vite", "^5.0.8", "Bundler y servidor de desarrollo."],
            ["@vitejs/plugin-react", "^4.2.1", "Plugin React para Vite."],
        ],
        caption="Tabla XIV. Dependencias del frontend",
    )

    add_heading(doc, "Apéndice C. Ejemplos de payloads JSON")
    add_paragraph(doc, "create_room (cliente → servidor):", bold=True)
    add_code_block(doc, '{ "username": "Ana" }')
    add_paragraph(doc, "room_created (servidor → cliente):", bold=True)
    add_code_block(doc, '{ "roomId": "ABC123", "state": { "status": "LOBBY", "players": [...] } }')
    add_paragraph(doc, "play_turn (cliente → servidor):", bold=True)
    add_code_block(doc, '{ "roomId": "ABC123", "cardIds": ["K♠_x7f2a", "K♥_b3c9"] }')
    add_paragraph(doc, "game_over (servidor → cliente):", bold=True)
    add_code_block(doc,
        '{ "loserId": "socket_xyz", "loserName": "Bob", '
        '"savedPlayers": ["Ana", "Carlos"], "reason": "El último jugador con cartas es el idiota." }'
    )

    add_heading(doc, "Apéndice D. Guía rápida para desarrolladores")
    add_bullet(doc, "Clonar el repositorio y ejecutar npm install en backend/ y frontend/.")
    add_bullet(doc, "Iniciar backend con npm start (puerto 3000).")
    add_bullet(doc, "Iniciar frontend con npm run dev (puerto 5173).")
    add_bullet(doc, "Abrir http://localhost:5173 en el navegador.")
    add_bullet(doc, "Constantes de eventos compartidas: frontend/src/constants/socketEvents.js.")
    add_bullet(doc, "Reglas de negocio autoritativas: backend/src/models/Game.js.")
    add_bullet(doc, "Health check del servidor: GET http://localhost:3000/health.")


def main():
    doc = Document()

    # Márgenes IEEE aproximados (1 pulgada)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    build_cover(doc)
    build_revision_table(doc)

    add_heading(doc, "Tabla de contenidos")
    add_toc(doc)
    doc.add_page_break()

    build_document(doc)

    os.makedirs(OUTPUT.parent, exist_ok=True)
    doc.save(str(OUTPUT))
    print(f"Documento generado: {OUTPUT}")
    print(f"Tablas: {_table_counter} | Figuras: {_figure_counter}")


if __name__ == "__main__":
    main()
